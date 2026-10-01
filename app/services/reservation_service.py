"""
Servicio de Cruce y Saldo de Reservas para Xdock (Logix).

Calcula el saldo pendiente del Reservation Log descontando lo ya despachado (AURRSGLBD0190)
para ser 100% precisos con las cantidades a reservar de Xdock en Inbound.

Regla:
1. Agrupa Despatched Qty por pedido e ítem.
2. Descuenta esa cantidad de las reservas del mismo pedido e ítem en orden FIFO/secuencial.
3. Si el saldo de una línea queda en cero, la elimina del resultado operativo.
4. Si queda un saldo positivo, conserva la línea y actualiza Action_QTY con el saldo pendiente.
5. Genera el archivo operativo exacto en RESERVATION_CSV_PATH y el libro de auditoría en Excel.
"""

from __future__ import annotations
import os
import re
import shutil
import time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

import polars as pl
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.core.config import (
    DATABASE_FOLDER,
    RESERVATION_CSV_PATH,
    RESERVATION_RAW_CSV_PATH,
    RESERVATIONS_AUDIT_EXCEL_PATH,
    get_despatched_file_path,
)


def normalize_identifier(value: Any) -> str:
    """Normaliza identificadores (SO Number, Item Code, Order Number).
    Convierte 00123, 123 y 123.0 al mismo valor numérico estándar.
    """
    if value is None:
        return ""
    text = re.sub(r"\s+", " ", str(value).strip().upper())
    if re.fullmatch(r"[+-]?\d+(?:\.0+)?", text):
        try:
            text = str(int(Decimal(text)))
        except Exception:
            pass
    return text


def parse_quantity(value: Any) -> float:
    """Parsea cantidades numéricas admitiendo formatos con comas y puntos."""
    if value is None:
        return 0.0
    text = str(value).strip()
    if not text:
        return 0.0

    text = text.replace(" ", "")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        parts = text.split(",")
        if len(parts[-1]) != 3:
            text = text.replace(",", ".")
        else:
            text = text.replace(",", "")

    try:
        val = float(text)
        return max(0.0, val)
    except (ValueError, InvalidOperation):
        return 0.0


def read_file_polars(file_path: Union[str, Path]) -> pl.DataFrame:
    """Lee un archivo CSV o Excel usando Polars."""
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix in {".xlsx", ".xlsm", ".xls"}:
        return pl.read_excel(path)
    elif suffix in {".csv", ".txt"}:
        for encoding in ("utf-8-sig", "utf-8", "latin-1"):
            try:
                return pl.read_csv(
                    path,
                    infer_schema_length=0,
                    encoding=encoding,
                    ignore_errors=True,
                )
            except Exception:
                continue
        # Fallback estándar
        return pl.read_csv(path, infer_schema_length=0, ignore_errors=True)
    else:
        raise ValueError(f"Formato no soportado para lectura: {suffix}")


def find_column(df_columns: List[str], target_candidates: List[str]) -> Optional[str]:
    """Encuentra el nombre exacto de la columna en base a una lista de nombres candidatos."""
    clean_map = {re.sub(r"[\s_]+", "", col.lower()): col for col in df_columns}
    for candidate in target_candidates:
        candidate_clean = re.sub(r"[\s_]+", "", candidate.lower())
        if candidate_clean in clean_map:
            return clean_map[candidate_clean]
    return None


def calculate_pending_reservations(
    reservation_path: Union[str, Path],
    dispatched_path: Union[str, Path],
    output_operational_csv: Optional[Union[str, Path]] = None,
    output_audit_excel: Optional[Union[str, Path]] = None,
) -> Tuple[pl.DataFrame, Dict[str, Any]]:
    """Ejecuta el algoritmo de cruce secuencial de saldos entre Reservas y Despachos.

    Retorna el DataFrame con las reservas operativas vigentes y un diccionario de métricas.
    """
    t0 = time.time()

    # 1. Leer Despachos
    df_dsp = read_file_polars(dispatched_path)
    dsp_cols = df_dsp.columns

    order_dsp_col = find_column(dsp_cols, ["Order Number", "Order_Number", "Order", "SO Number", "SO_Number"])
    item_dsp_col = find_column(dsp_cols, ["Item Code", "Item_Code", "Item", "Codigo"])
    qty_dsp_col = find_column(dsp_cols, ["Despatched Qty", "Despatched_Qty", "Despatch Qty", "Qty", "Cantidad"])

    if not order_dsp_col or not item_dsp_col or not qty_dsp_col:
        raise KeyError(
            f"En Despachos faltan columnas requeridas. Disponibles: {dsp_cols}. "
            f"Detectadas: Order={order_dsp_col}, Item={item_dsp_col}, Qty={qty_dsp_col}"
        )

    # Agrupar despachos por (Order, Item)
    dsp_records = df_dsp.select([order_dsp_col, item_dsp_col, qty_dsp_col]).iter_rows(named=True)
    dispatched_totals: Dict[Tuple[str, str], float] = {}

    for r in dsp_records:
        ord_key = normalize_identifier(r[order_dsp_col])
        itm_key = normalize_identifier(r[item_dsp_col])
        if ord_key and itm_key:
            qty_val = parse_quantity(r[qty_dsp_col])
            if qty_val > 0:
                key = (ord_key, itm_key)
                dispatched_totals[key] = dispatched_totals.get(key, 0.0) + qty_val

    # 2. Leer Reservas
    df_res = read_file_polars(reservation_path)
    res_cols = df_res.columns

    so_res_col = find_column(res_cols, ["SO_Number", "SO Number", "SO_Num", "SO", "Order_Number", "Order Number"])
    item_res_col = find_column(res_cols, ["Item_Code", "Item Code", "Item", "Codigo"])
    qty_res_col = find_column(res_cols, ["Action_QTY", "Action Qty", "Action_Qty", "Quantity_reserved", "Qty"])
    line_res_col = find_column(res_cols, ["SO_Line_Number", "SO Line Number", "SO_Line", "SO Line", "Line_Number", "Line Number", "Line"])
    po_res_col = find_column(res_cols, ["PO_Number", "PO Number", "PO_Num", "PO", "Customer_Reference", "Customer Reference"])
    cust_res_col = find_column(res_cols, ["Customer_Name", "Customer Name", "Customer_Code", "Customer Code", "Customer"])

    if not so_res_col or not item_res_col or not qty_res_col:
        raise KeyError(
            f"En Reservas faltan columnas requeridas. Disponibles: {res_cols}. "
            f"Detectadas: SO={so_res_col}, Item={item_res_col}, Qty={qty_res_col}"
        )

    # Desduplicación preventiva de líneas de reserva (idéntico a Rust Core)
    so_line_map: Dict[Tuple[str, str, str, str, str], Dict[str, Any]] = {}
    for row in df_res.iter_rows(named=True):
        ord_key = normalize_identifier(row[so_res_col])
        line_key = normalize_identifier(row[line_res_col]) if line_res_col else ""
        itm_key = normalize_identifier(row[item_res_col])
        po_key = normalize_identifier(row[po_res_col]) if po_res_col else ""
        cust_key = normalize_identifier(row[cust_res_col]) if cust_res_col else ""
        qty_val = parse_quantity(row[qty_res_col])

        if ord_key and line_key and itm_key:
            dedup_key = (ord_key, line_key, itm_key, po_key, cust_key)
            if dedup_key not in so_line_map:
                so_line_map[dedup_key] = row
            else:
                existing_qty = parse_quantity(so_line_map[dedup_key][qty_res_col])
                if qty_val > existing_qty:
                    so_line_map[dedup_key] = row

    # Mantener el orden original de primera aparición al reconstruir la lista desduplicada
    final_input_rows: List[Dict[str, Any]] = []
    seen_keys = set()
    for row in df_res.iter_rows(named=True):
        ord_key = normalize_identifier(row[so_res_col])
        line_key = normalize_identifier(row[line_res_col]) if line_res_col else ""
        itm_key = normalize_identifier(row[item_res_col])
        po_key = normalize_identifier(row[po_res_col]) if po_res_col else ""
        cust_key = normalize_identifier(row[cust_res_col]) if cust_res_col else ""

        if ord_key and line_key and itm_key:
            dedup_key = (ord_key, line_key, itm_key, po_key, cust_key)
            if dedup_key not in seen_keys:
                seen_keys.add(dedup_key)
                final_input_rows.append(so_line_map[dedup_key])
        else:
            final_input_rows.append(row)

    dedup_eliminated = df_res.height - len(final_input_rows)
    remaining_dispatch = dispatched_totals.copy()

    operational_rows: List[Dict[str, Any]] = []
    audit_rows: List[Dict[str, Any]] = []
    eliminated_count = 0

    for row in final_input_rows:
        ord_key = normalize_identifier(row[so_res_col])
        itm_key = normalize_identifier(row[item_res_col])
        orig_qty = parse_quantity(row[qty_res_col])

        key = (ord_key, itm_key)
        available_dsp = remaining_dispatch.get(key, 0.0) if (ord_key and itm_key) else 0.0
        applied_qty = min(orig_qty, available_dsp)
        residual_qty = max(0.0, orig_qty - applied_qty)
        if ord_key and itm_key:
            remaining_dispatch[key] = available_dsp - applied_qty

        # Formato numérico limpio para la salida
        formatted_residual = (
            str(int(residual_qty)) if residual_qty == int(residual_qty) else f"{residual_qty:.4f}".rstrip("0").rstrip(".")
        )

        # Registro operativo si tiene saldo > 0
        if residual_qty > 0.0001:
            op_row = dict(row)
            op_row[qty_res_col] = formatted_residual
            operational_rows.append(op_row)
        else:
            eliminated_count += 1

        # Auditoría completa
        aud_row = dict(row)
        aud_row["Cantidad Reserva Original"] = int(orig_qty) if orig_qty == int(orig_qty) else orig_qty
        aud_row["Cantidad Despachada Aplicada"] = int(applied_qty) if applied_qty == int(applied_qty) else applied_qty
        aud_row["Saldo Pendiente"] = int(residual_qty) if residual_qty == int(residual_qty) else residual_qty
        aud_row["Estado Cruce"] = "ELIMINADO - SALDO CERO" if residual_qty <= 0.0001 else "CONSERVADO - SALDO PENDIENTE"
        audit_rows.append(aud_row)

    # DataFrame operativo final conservando el orden de columnas original
    df_operational = pl.DataFrame(operational_rows, schema=df_res.schema) if operational_rows else pl.DataFrame([], schema=df_res.schema)

    # Guardar CSV operativo si se solicitó
    if output_operational_csv:
        Path(output_operational_csv).parent.mkdir(parents=True, exist_ok=True)
        df_operational.write_csv(output_operational_csv)

    # 3. Despachos en exceso o sin reserva
    excess_rows = []
    for (ord_k, itm_k), rem_qty in remaining_dispatch.items():
        if rem_qty > 0.0001:
            excess_rows.append({
                "Order Number": ord_k,
                "Item Code": itm_k,
                "Despatched Qty sin aplicar": int(rem_qty) if rem_qty == int(rem_qty) else rem_qty,
                "Observacion": "Despacho superior a la reserva o sin reserva coincidente",
            })

    # 4. Generar Libro de Auditoría Excel si se especificó
    if output_audit_excel:
        Path(output_audit_excel).parent.mkdir(parents=True, exist_ok=True)
        _build_audit_excel(audit_rows, excess_rows, Path(output_audit_excel))

    elapsed = round(time.time() - t0, 3)
    metrics = {
        "reservations_original": df_res.height,
        "reservations_deduplicated": len(final_input_rows),
        "reservations_duplicates_removed": dedup_eliminated,
        "reservations_operational": df_operational.height,
        "reservations_eliminated": eliminated_count,
        "dispatch_keys_total": len(dispatched_totals),
        "dispatch_excess_keys": len(excess_rows),
        "elapsed_seconds": elapsed,
        "operational_csv": str(output_operational_csv) if output_operational_csv else "",
        "audit_excel": str(output_audit_excel) if output_audit_excel else "",
    }

    print(
        f"[XDOCK CRUCE] Procesado en {elapsed}s: {df_res.height} reservas brutas ({dedup_eliminated} duplicados eliminados) -> "
        f"{df_operational.height} con saldo pendiente, {eliminated_count} eliminadas (saldo 0).",
        flush=True,
    )
    return df_operational, metrics
    return df_operational, metrics


import xlsxwriter

def _build_audit_excel(audit_rows: List[Dict[str, Any]], excess_rows: List[Dict[str, Any]], output_path: Path) -> None:
    """Construye un Excel de auditoría estilizado con 4 hojas: Control completo, Saldos pendientes, Saldos cero y Despachos sin aplicar usando xlsxwriter para velocidad extrema."""
    wb = xlsxwriter.Workbook(str(output_path), {"constant_memory": True})

    header_fmt = wb.add_format({
        "bold": True,
        "bg_color": "#1F4E79",
        "font_color": "#FFFFFF",
        "align": "center",
        "valign": "vcenter",
        "font_name": "Calibri",
        "font_size": 10,
    })
    cell_fmt = wb.add_format({
        "font_name": "Calibri",
        "font_size": 9,
    })

    headers = list(audit_rows[0].keys()) if audit_rows else []

    def populate_sheet(sheet_name: str, rows: List[Dict[str, Any]], custom_headers: Optional[List[str]] = None):
        ws = wb.add_worksheet(sheet_name)
        cols = custom_headers or headers
        if not cols:
            return
        ws.write_row(0, 0, cols, header_fmt)
        ws.freeze_panes(1, 0)

        # Estimar anchos basados en nombres de columna y muestra inicial
        col_widths = {i: max(len(str(c)), 10) for i, c in enumerate(cols)}
        sample_rows = rows[:50]
        for r_data in sample_rows:
            for i, c in enumerate(cols):
                val_len = len(str(r_data.get(c, "") or ""))
                if val_len > col_widths[i]:
                    col_widths[i] = min(val_len + 2, 38)

        for i, w in col_widths.items():
            ws.set_column(i, i, w)

        for row_idx, r in enumerate(rows, start=1):
            row_vals = [r.get(c, "") for c in cols]
            ws.write_row(row_idx, 0, row_vals, cell_fmt)

        if rows:
            ws.autofilter(0, 0, len(rows), len(cols) - 1)

    pending_rows = [r for r in audit_rows if r.get("Saldo Pendiente", 0) > 0.0001]
    zero_rows = [r for r in audit_rows if r.get("Saldo Pendiente", 0) <= 0.0001]

    populate_sheet("Control completo", audit_rows)
    populate_sheet("Saldos pendientes", pending_rows)
    populate_sheet("Saldos cero", zero_rows)

    excess_headers = ["Order Number", "Item Code", "Despatched Qty sin aplicar", "Observacion"]
    populate_sheet("Despacho sin aplicar", excess_rows, excess_headers)

    wb.close()


async def recalculate_xdock_reservations_task(generate_cache: bool = True) -> Dict[str, Any]:
    """Tarea asíncrona de alto nivel para ejecutar el cruce y recargar el caché de Xdock.

    1. Verifica si existe RESERVATION_RAW_CSV_PATH (o crea copia desde RESERVATION_CSV_PATH).
    2. Busca si existe archivo de despachos (AURRSGLBD0190 .xlsx o .csv).
    3. Si ambos existen, ejecuta calculate_pending_reservations.
    4. Si no hay archivo de despachos, conserva las reservas brutas.
    5. Recarga el caché de memoria en csv_handler.
    """
    from app.services import csv_handler

    # Asegurar archivo de reservas crudo
    if not os.path.exists(RESERVATION_RAW_CSV_PATH):
        if os.path.exists(RESERVATION_CSV_PATH):
            shutil.copyfile(RESERVATION_CSV_PATH, RESERVATION_RAW_CSV_PATH)
        else:
            return {"status": "skipped", "message": "No existe archivo de reservas."}

    dsp_path = get_despatched_file_path()

    if dsp_path and os.path.exists(dsp_path):
        try:
            _, metrics = calculate_pending_reservations(
                reservation_path=RESERVATION_RAW_CSV_PATH,
                dispatched_path=dsp_path,
                output_operational_csv=RESERVATION_CSV_PATH,
                output_audit_excel=RESERVATIONS_AUDIT_EXCEL_PATH,
            )
            if generate_cache:
                await csv_handler.generate_reservation_cache()
            return {"status": "success", "cruce": True, "metrics": metrics}
        except Exception as e:
            print(f"[ERROR CRUCE XDOCK] Error calculando saldos: {e}", flush=True)
            # Fallback: mantener el archivo crudo
            shutil.copyfile(RESERVATION_RAW_CSV_PATH, RESERVATION_CSV_PATH)
            if generate_cache:
                await csv_handler.generate_reservation_cache()
            return {"status": "error", "error": str(e)}
    else:
        # No hay archivo de despachos: mantener reservas directas
        if os.path.exists(RESERVATION_RAW_CSV_PATH):
            shutil.copyfile(RESERVATION_RAW_CSV_PATH, RESERVATION_CSV_PATH)
        if generate_cache:
            await csv_handler.generate_reservation_cache()
        return {
            "status": "success",
            "cruce": False,
            "message": "Archivo de reservas cargado sin cruce (no se encontró AURRSGLBD0190 de despachos).",
        }


def main() -> int:
    """Punto de entrada para ejecución directa desde línea de comandos (CLI)."""
    import argparse
    import sys

    parser = argparse.ArgumentParser(
        description="Descontar cantidades despachadas (0190) y conservar saldos pendientes de Xdock."
    )
    parser.add_argument(
        "reservation_file",
        nargs="?",
        type=Path,
        default=None,
        help="Ruta al archivo de reservas (CSV). Si se omite, usa databases/AURRSLAMP0006_raw.csv o AURRSLAMP0006.csv",
    )
    parser.add_argument(
        "despatched_file",
        nargs="?",
        type=Path,
        default=None,
        help="Ruta al archivo de despachos (XLSX o CSV). Si se omite, busca en databases/AURRSGLBD0190*",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Ruta de salida del CSV operativo filtrado.",
    )
    parser.add_argument(
        "--audit",
        type=Path,
        default=None,
        help="Ruta de salida del Excel de auditoría completa.",
    )
    args = parser.parse_args()

    res_path = args.reservation_file
    if not res_path:
        raw_p = Path(RESERVATION_RAW_CSV_PATH)
        std_p = Path(RESERVATION_CSV_PATH)
        res_path = raw_p if raw_p.exists() else std_p

    dsp_path = args.despatched_file
    if not dsp_path:
        found_dsp = get_despatched_file_path()
        if found_dsp:
            dsp_path = Path(found_dsp)

    if not res_path or not res_path.exists():
        print(f"ERROR: No se encontró el archivo de reservas: {res_path}", file=sys.stderr)
        return 1

    if not dsp_path or not dsp_path.exists():
        print(f"ERROR: No se encontró el archivo de despachos: {dsp_path}", file=sys.stderr)
        return 1

    output_csv = args.output or Path(RESERVATION_CSV_PATH)
    output_audit = args.audit or Path(RESERVATIONS_AUDIT_EXCEL_PATH)

    print(f"-> Archivo Reservas:   {res_path}")
    print(f"-> Archivo Despachos:  {dsp_path}")
    print(f"-> Salida Operativa:   {output_csv}")
    print(f"-> Salida Auditoría:   {output_audit}")
    print("Procesando cruce de saldos con Polars...")

    df_op, metrics = calculate_pending_reservations(
        reservation_path=res_path,
        dispatched_path=dsp_path,
        output_operational_csv=output_csv,
        output_audit_excel=output_audit,
    )

    print("\n--- RESUMEN DEL CRUCE ---")
    print(f"Líneas de reserva originales:    {metrics['reservations_original']:,}")
    print(f"Líneas con saldo pendiente:      {metrics['reservations_operational']:,}")
    print(f"Líneas eliminadas (saldo cero):  {metrics['reservations_eliminated']:,}")
    print(f"Claves de despacho agrupadas:    {metrics['dispatch_keys_total']:,}")
    print(f"Tiempo de ejecución:             {metrics['elapsed_seconds']} s")
    print(f"Resultado guardado en:           {output_csv.resolve()}")
    print(f"Auditoría guardada en:           {output_audit.resolve()}")
    return 0


if __name__ == "__main__":
    import sys
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        import traceback
        traceback.print_exc()
        raise SystemExit(1)

